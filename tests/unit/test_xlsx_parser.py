"""Unit tests for the structure-preserving XLSX parser."""

from pathlib import Path

import pytest
from openpyxl import Workbook

from app.ingestion.schemas import SpreadsheetBlock
from app.ingestion.xlsx_parser import XlsxParser


def _save_workbook(path: Path) -> None:
    workbook = Workbook()

    first = workbook.active
    first.title = "Decisions"
    first["A1"] = "Decision"
    first["B1"] = "Amount"
    first["A2"] = "Standard contract"
    first["B2"] = 100000
    first["C2"] = "=B2*2"

    second = workbook.create_sheet("Control")
    second["A1"] = "Document ID"
    second["B1"] = "DOC-TEST"

    workbook.save(path)
    workbook.close()


def test_parser_preserves_worksheet_boundaries(tmp_path: Path) -> None:
    path = tmp_path / "sample.xlsx"
    _save_workbook(path)

    content = XlsxParser().parse(path)

    assert len(content.blocks) == 2
    assert all(isinstance(block, SpreadsheetBlock) for block in content.blocks)
    assert [block.sheet_name for block in content.blocks] == [
        "Decisions",
        "Control",
    ]


def test_parser_preserves_sheet_ranges(tmp_path: Path) -> None:
    path = tmp_path / "sample.xlsx"
    _save_workbook(path)

    content = XlsxParser().parse(path)

    assert content.blocks[0].cell_range == "A1:C2"
    assert content.blocks[1].cell_range == "A1:B1"


def test_parser_preserves_cell_coordinates_values_and_types(
    tmp_path: Path,
) -> None:
    path = tmp_path / "sample.xlsx"
    _save_workbook(path)

    content = XlsxParser().parse(path)
    cells = {cell.coordinate: cell for cell in content.blocks[0].cells}

    assert cells["A1"].value == "Decision"
    assert cells["A1"].data_type == "s"
    assert cells["B2"].value == 100000
    assert cells["B2"].data_type == "n"


def test_parser_preserves_formula_in_formula_field(tmp_path: Path) -> None:
    path = tmp_path / "sample.xlsx"
    _save_workbook(path)

    content = XlsxParser().parse(path)
    cells = {cell.coordinate: cell for cell in content.blocks[0].cells}

    assert cells["C2"].formula == "=B2*2"
    assert cells["C2"].data_type == "f"
    assert cells["C2"].value is None


def test_parser_omits_empty_cells(tmp_path: Path) -> None:
    path = tmp_path / "sample.xlsx"
    _save_workbook(path)

    content = XlsxParser().parse(path)

    coordinates = {
        cell.coordinate
        for block in content.blocks
        for cell in block.cells
    }

    assert "C1" not in coordinates


def test_parser_assigns_stable_block_order(tmp_path: Path) -> None:
    path = tmp_path / "sample.xlsx"
    _save_workbook(path)

    content = XlsxParser().parse(path)

    assert [block.order for block in content.blocks] == [0, 1]
    assert [block.block_id for block in content.blocks] == [
        "xlsx-sheet-1",
        "xlsx-sheet-2",
    ]


def test_parser_rejects_missing_file(tmp_path: Path) -> None:
    path = tmp_path / "missing.xlsx"

    with pytest.raises(FileNotFoundError):
        XlsxParser().parse(path)


def test_parser_rejects_non_xlsx_source(tmp_path: Path) -> None:
    path = tmp_path / "sample.txt"
    path.write_text("not an xlsx workbook", encoding="utf-8")

    with pytest.raises(ValueError, match="Unsupported XLSX source"):
        XlsxParser().parse(path)


def test_parser_rejects_workbook_without_content(tmp_path: Path) -> None:
    path = tmp_path / "empty.xlsx"
    workbook = Workbook()
    workbook.save(path)
    workbook.close()

    with pytest.raises(ValueError, match="No spreadsheet content extracted"):
        XlsxParser().parse(path)
