"""Unit tests for deterministic native PPTX chart enrichment."""

from __future__ import annotations

import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from app.ingestion.pptx_native_enricher import PptxNativeEnricher

_CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def _add_chart(
    slide,
    *,
    chart_type=XL_CHART_TYPE.COLUMN_CLUSTERED,
    categories=("Period 1", "Period 2", "Period 3"),
    series=(("Forecast", (1.0, 2.0, 3.0)),),
    title: str | None = None,
):
    chart_data = ChartData()
    chart_data.categories = categories
    for name, values in series:
        chart_data.add_series(name, values)

    chart = slide.shapes.add_chart(
        chart_type,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        chart_data,
    ).chart

    if title is not None:
        chart.has_title = True
        chart.chart_title.text_frame.text = title

    return chart


def _save_presentation(path: Path, build) -> Path:
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    build(presentation)
    presentation.save(path)
    return path


def _rewrite_first_formula_sheet(
    source: Path,
    target: Path,
    *,
    sheet_name: str,
) -> Path:
    """Rewrite one chart formula to a missing sheet without touching the workbook."""
    with zipfile.ZipFile(source, "r") as source_zip:
        members = {name: source_zip.read(name) for name in source_zip.namelist()}

    chart_name = sorted(
        name
        for name in members
        if name.startswith("ppt/charts/chart") and name.endswith(".xml")
    )[0]
    root = ET.fromstring(members[chart_name])
    formula = root.find(
        f".//{{{_CHART_NS}}}ser/{{{_CHART_NS}}}cat//{{{_CHART_NS}}}f"
    )
    if formula is None:
        formula = root.find(
            f".//{{{_CHART_NS}}}ser/{{{_CHART_NS}}}val//{{{_CHART_NS}}}f"
        )

    assert formula is not None
    assert formula.text is not None

    _, cell_range = formula.text.rsplit("!", 1)
    formula.text = f"'{sheet_name}'!{cell_range}"
    members[chart_name] = ET.tostring(
        root,
        encoding="utf-8",
        xml_declaration=True,
    )

    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as target_zip:
        for name, payload in members.items():
            target_zip.writestr(name, payload)

    return target


def test_extract_charts_preserves_native_chart_structure(tmp_path: Path) -> None:
    source = _save_presentation(
        tmp_path / "chart.pptx",
        lambda presentation: _add_chart(
            presentation.slides[0],
            categories=("Period 1", "Period 2", "Total"),
            series=(
                ("Actual", (-0.07, 0.04, 0.12)),
                ("Forecast", (0.02, 0.03, 0.05)),
            ),
            title="Actual vs. Forecast",
        ),
    )

    blocks = PptxNativeEnricher().extract_charts(source)

    assert len(blocks) == 1
    chart = blocks[0]
    assert chart.block_id == "pptx-chart-s001-001"
    assert chart.order == 0
    assert chart.chart_type == "COLUMN_CLUSTERED"
    assert chart.title == "Actual vs. Forecast"
    assert chart.categories == ("Period 1", "Period 2", "Total")
    assert tuple(series.name for series in chart.series) == ("Actual", "Forecast")
    assert chart.series[0].values == (-0.07, 0.04, 0.12)
    assert chart.location.slide_number == 1
    assert chart.location.page_number is None


def test_extract_charts_preserves_slide_provenance(tmp_path: Path) -> None:
    def build(presentation):
        second = presentation.slides.add_slide(presentation.slide_layouts[6])
        _add_chart(second, chart_type=XL_CHART_TYPE.LINE)

    source = _save_presentation(tmp_path / "slides.pptx", build)

    blocks = PptxNativeEnricher().extract_charts(source)

    assert len(blocks) == 1
    assert blocks[0].location.slide_number == 2
    assert blocks[0].block_id == "pptx-chart-s002-001"


def test_extract_charts_assigns_deterministic_ids_and_global_order(
    tmp_path: Path,
) -> None:
    def build(presentation):
        first = presentation.slides[0]
        _add_chart(first)
        _add_chart(first, chart_type=XL_CHART_TYPE.LINE)

        second = presentation.slides.add_slide(presentation.slide_layouts[6])
        _add_chart(second)

    source = _save_presentation(tmp_path / "multiple.pptx", build)

    blocks = PptxNativeEnricher().extract_charts(source)

    assert tuple(block.block_id for block in blocks) == (
        "pptx-chart-s001-001",
        "pptx-chart-s001-002",
        "pptx-chart-s002-001",
    )
    assert tuple(block.order for block in blocks) == (0, 1, 2)


def test_extract_charts_preserves_chart_without_title(tmp_path: Path) -> None:
    source = _save_presentation(
        tmp_path / "untitled.pptx",
        lambda presentation: _add_chart(presentation.slides[0]),
    )

    assert PptxNativeEnricher().extract_charts(source)[0].title is None


def test_extract_charts_returns_empty_tuple_when_no_charts(tmp_path: Path) -> None:
    source = _save_presentation(tmp_path / "no-charts.pptx", lambda _: None)

    assert PptxNativeEnricher().extract_charts(source) == ()


def test_extract_charts_resolves_ooxml_lineage_to_embedded_workbook(
    tmp_path: Path,
) -> None:
    source = _save_presentation(
        tmp_path / "resolved.pptx",
        lambda presentation: _add_chart(
            presentation.slides[0],
            categories=("A", "B", "C"),
            series=(
                ("Series A", (10.0, 20.0, 30.0)),
                ("Series B", (11.0, 21.0, 31.0)),
            ),
        ),
    )

    chart = PptxNativeEnricher().extract_charts(source)[0]

    category_ref = chart.category_source_reference
    assert category_ref is not None
    assert category_ref.status == "resolved"
    assert category_ref.resource.startswith("ppt/embeddings/")
    assert category_ref.resource.endswith(".xlsx")
    assert category_ref.sheet_name is not None
    assert category_ref.cell_range is not None

    assert len(chart.series) == 2
    for series in chart.series:
        assert series.source_reference is not None
        assert series.source_reference.status == "resolved"
        assert series.source_reference.resource == category_ref.resource
        assert series.source_reference.sheet_name == category_ref.sheet_name
        assert series.source_reference.cell_range is not None


def test_extract_charts_preserves_unresolved_ooxml_reference(
    tmp_path: Path,
) -> None:
    original = _save_presentation(
        tmp_path / "original.pptx",
        lambda presentation: _add_chart(presentation.slides[0]),
    )
    source = _rewrite_first_formula_sheet(
        original,
        tmp_path / "unresolved.pptx",
        sheet_name="MissingSheet",
    )

    chart = PptxNativeEnricher().extract_charts(source)[0]

    references = [
        chart.category_source_reference,
        *(series.source_reference for series in chart.series),
    ]
    unresolved = [
        reference
        for reference in references
        if reference is not None and reference.status == "unresolved"
    ]

    assert len(unresolved) == 1
    assert unresolved[0].sheet_name == "MissingSheet"
    assert unresolved[0].cell_range is not None
    assert unresolved[0].resource.startswith("ppt/embeddings/")


def test_extract_charts_does_not_materialize_embedded_workbook_as_blocks(
    tmp_path: Path,
) -> None:
    source = _save_presentation(
        tmp_path / "lineage-only.pptx",
        lambda presentation: _add_chart(
            presentation.slides[0],
            series=(("Series A", (1.0, 2.0, 3.0)),),
        ),
    )

    blocks = PptxNativeEnricher().extract_charts(source)

    assert len(blocks) == 1
    assert blocks[0].block_id.startswith("pptx-chart-")
    assert blocks[0].series[0].source_reference is not None


def test_extract_charts_rejects_missing_file(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        PptxNativeEnricher().extract_charts(tmp_path / "missing.pptx")


def test_extract_charts_rejects_non_pptx_file(tmp_path: Path) -> None:
    source = tmp_path / "not-pptx.txt"
    source.write_text("not a presentation", encoding="utf-8")

    with pytest.raises(ValueError, match="only supports .pptx"):
        PptxNativeEnricher().extract_charts(source)


def test_extract_charts_requires_path_instance() -> None:
    with pytest.raises(TypeError, match="pathlib.Path"):
        PptxNativeEnricher().extract_charts(
            "presentation.pptx"  # type: ignore[arg-type]
        )
