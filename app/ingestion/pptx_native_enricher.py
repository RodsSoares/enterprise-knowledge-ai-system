"""Deterministic native PPTX enrichment for presentation-specific structures.

The general document structure remains the responsibility of DoclingParser.
This module adds presentation-native chart structure and deterministic OOXML
lineage without exposing python-pptx or OOXML models downstream.
"""

from __future__ import annotations

import io
import posixpath
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from pptx import Presentation

from app.ingestion.schemas import (
    ChartBlock,
    ChartSeries,
    ChartSourceReference,
    SourceLocation,
)

_CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_PACKAGE_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/package"
)
_EXCEL_REFERENCE = re.compile(
    r"^(?P<sheet>'(?:[^']|'')+'|[^!]+)!(?P<range>\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?)$"
)


class PptxNativeEnricher:
    """Extract native PPTX charts into application-controlled canonical blocks."""

    def extract_charts(self, source_path: Path) -> tuple[ChartBlock, ...]:
        """Return native charts in deterministic slide/shape order."""
        self._validate_source(source_path)

        presentation = Presentation(source_path)
        blocks: list[ChartBlock] = []
        order = 0

        with zipfile.ZipFile(source_path) as package:
            for slide_number, slide in enumerate(presentation.slides, start=1):
                chart_number = 0

                for shape in slide.shapes:
                    if not getattr(shape, "has_chart", False):
                        continue

                    chart_number += 1
                    chart = shape.chart
                    chart_part = str(chart.part.partname).lstrip("/")
                    lineage = self._chart_lineage(package, chart_part)

                    native_series = tuple(chart.series)
                    canonical_series: list[ChartSeries] = []
                    for index, series in enumerate(native_series):
                        source_reference = (
                            lineage["series"][index]
                            if index < len(lineage["series"])
                            else None
                        )
                        canonical_series.append(
                            ChartSeries(
                                name=self._series_name(series),
                                values=tuple(getattr(series, "values", ()) or ()),
                                source_reference=source_reference,
                            )
                        )

                    if not canonical_series:
                        raise ValueError(
                            "Native PPTX chart contains no extractable series."
                        )

                    blocks.append(
                        ChartBlock(
                            block_id=(
                                f"pptx-chart-s{slide_number:03d}-{chart_number:03d}"
                            ),
                            order=order,
                            chart_type=self._chart_type_name(chart.chart_type),
                            title=self._chart_title(chart),
                            categories=self._chart_categories(chart),
                            category_source_reference=lineage["category"],
                            series=tuple(canonical_series),
                            location=SourceLocation(slide_number=slide_number),
                        )
                    )
                    order += 1

        return tuple(blocks)

    @staticmethod
    def _validate_source(source_path: Path) -> None:
        if not isinstance(source_path, Path):
            raise TypeError("source_path must be a pathlib.Path.")
        if not source_path.is_file():
            raise FileNotFoundError(source_path)
        if source_path.suffix.lower() != ".pptx":
            raise ValueError("PptxNativeEnricher only supports .pptx files.")

    @staticmethod
    def _chart_type_name(chart_type: object) -> str:
        name = getattr(chart_type, "name", None)
        if isinstance(name, str) and name.strip():
            return name

        value = str(chart_type).strip()
        if not value:
            raise ValueError("Native PPTX chart type could not be determined.")
        return value.split(" (", 1)[0]

    @staticmethod
    def _chart_title(chart) -> str | None:
        if not chart.has_title:
            return None
        text = chart.chart_title.text_frame.text.strip()
        return text or None

    @staticmethod
    def _chart_categories(chart) -> tuple[str | int | float | bool | None, ...]:
        for plot in chart.plots:
            categories = getattr(plot, "categories", None)
            if categories is None:
                continue
            values = tuple(getattr(category, "label", None) for category in categories)
            if values:
                return values
        return ()

    @staticmethod
    def _series_name(series) -> str | None:
        name = getattr(series, "name", None)
        if isinstance(name, str):
            return name.strip() or None
        return None

    def _chart_lineage(
        self,
        package: zipfile.ZipFile,
        chart_part: str,
    ) -> dict[str, object]:
        """Resolve OOXML chart formula references against its embedded workbook."""
        chart_xml = ET.fromstring(package.read(chart_part))
        workbook_part = self._embedded_workbook_part(package, chart_part)
        sheet_names = self._workbook_sheet_names(package, workbook_part)

        series_nodes = chart_xml.findall(f".//{{{_CHART_NS}}}ser")
        series_references: list[ChartSourceReference | None] = []
        category_reference: ChartSourceReference | None = None

        for series_node in series_nodes:
            if category_reference is None:
                category_formula = self._first_formula(
                    series_node,
                    ("cat", "xVal"),
                )
                category_reference = self._source_reference(
                    category_formula,
                    workbook_part,
                    sheet_names,
                )

            value_formula = self._first_formula(series_node, ("val", "yVal"))
            series_references.append(
                self._source_reference(
                    value_formula,
                    workbook_part,
                    sheet_names,
                )
            )

        return {
            "category": category_reference,
            "series": tuple(series_references),
        }

    @staticmethod
    def _first_formula(
        series_node: ET.Element,
        containers: tuple[str, ...],
    ) -> str | None:
        for container in containers:
            node = series_node.find(
                f"./{{{_CHART_NS}}}{container}//{{{_CHART_NS}}}f"
            )
            if node is not None and node.text and node.text.strip():
                return node.text.strip()
        return None

    @staticmethod
    def _embedded_workbook_part(
        package: zipfile.ZipFile,
        chart_part: str,
    ) -> str | None:
        chart_dir = posixpath.dirname(chart_part)
        rels_part = posixpath.join(
            chart_dir,
            "_rels",
            posixpath.basename(chart_part) + ".rels",
        )
        if rels_part not in package.namelist():
            return None

        root = ET.fromstring(package.read(rels_part))
        for relationship in root.findall(f"{{{_REL_NS}}}Relationship"):
            rel_type = relationship.attrib.get("Type", "")
            target = relationship.attrib.get("Target")
            target_mode = relationship.attrib.get("TargetMode", "")

            if not target or target_mode.lower() == "external":
                continue

            if (
                rel_type == _PACKAGE_REL_TYPE
                or rel_type.endswith("/package")
            ):
                return posixpath.normpath(
                    posixpath.join(chart_dir, target)
                )

        return None

    @staticmethod
    def _workbook_sheet_names(
        package: zipfile.ZipFile,
        workbook_part: str | None,
    ) -> frozenset[str]:
        if workbook_part is None or workbook_part not in package.namelist():
            return frozenset()

        workbook = load_workbook(
            io.BytesIO(package.read(workbook_part)),
            read_only=True,
            data_only=False,
        )
        try:
            return frozenset(workbook.sheetnames)
        finally:
            workbook.close()

    @staticmethod
    def _source_reference(
        formula: str | None,
        workbook_part: str | None,
        sheet_names: frozenset[str],
    ) -> ChartSourceReference | None:
        if formula is None:
            return None

        match = _EXCEL_REFERENCE.fullmatch(formula)
        if match is None:
            return ChartSourceReference(
                resource=workbook_part or "embedded-workbook",
                status="unresolved",
            )

        raw_sheet = match.group("sheet")
        if raw_sheet.startswith("'") and raw_sheet.endswith("'"):
            sheet_name = raw_sheet[1:-1].replace("''", "'")
        else:
            sheet_name = raw_sheet

        cell_range = match.group("range")
        resolved = workbook_part is not None and sheet_name in sheet_names

        return ChartSourceReference(
            resource=workbook_part or "embedded-workbook",
            sheet_name=sheet_name,
            cell_range=cell_range,
            status="resolved" if resolved else "unresolved",
        )
