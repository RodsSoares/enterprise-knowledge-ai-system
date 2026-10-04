"""Structure-preserving XLSX parser backed by openpyxl."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

from app.ingestion.schemas import (
    NormalizedContent,
    SpreadsheetBlock,
    SpreadsheetCell,
)


class XlsxParser:
    """Parse XLSX workbooks into the application's canonical content model.

    XLSX is handled with openpyxl rather than flattened through a generic
    document parser so worksheet boundaries, cell coordinates, values,
    formulas, and data types remain explicit.
    """

    def parse(self, source_path: Path) -> NormalizedContent:
        source_path = Path(source_path)

        if not source_path.exists():
            raise FileNotFoundError(source_path)
        if not source_path.is_file():
            raise ValueError(f"XLSX source must be a file: {source_path}")
        if source_path.suffix.lower() != ".xlsx":
            raise ValueError(f"Unsupported XLSX source: {source_path}")

        workbook = load_workbook(
            filename=source_path,
            data_only=False,
            read_only=False,
        )

        try:
            blocks = tuple(
                self._worksheet_block(worksheet, order)
                for order, worksheet in enumerate(workbook.worksheets)
                if self._has_content(worksheet)
            )
        finally:
            workbook.close()

        if not blocks:
            raise ValueError(f"No spreadsheet content extracted from: {source_path}")

        return NormalizedContent(blocks=blocks)

    @staticmethod
    def _has_content(worksheet: Worksheet) -> bool:
        return any(
            cell.value is not None
            for row in worksheet.iter_rows()
            for cell in row
        )

    def _worksheet_block(
        self,
        worksheet: Worksheet,
        order: int,
    ) -> SpreadsheetBlock:
        cells = tuple(
            self._spreadsheet_cell(cell)
            for row in worksheet.iter_rows()
            for cell in row
            if cell.value is not None
        )

        return SpreadsheetBlock(
            block_id=f"xlsx-sheet-{order + 1}",
            order=order,
            sheet_name=worksheet.title,
            cells=cells,
            cell_range=worksheet.calculate_dimension(),
        )

    @staticmethod
    def _spreadsheet_cell(cell: Cell) -> SpreadsheetCell:
        is_formula = cell.data_type == "f"

        return SpreadsheetCell(
            coordinate=cell.coordinate,
            value=None if is_formula else cell.value,
            formula=str(cell.value) if is_formula else None,
            data_type=cell.data_type,
        )
